import pandas as pd
from openpyxl import load_workbook

from src.risk_monitoring_excel import (
    create_risk_monitoring_excel,
)


def test_create_risk_monitoring_excel(tmp_path):
    dashboard = pd.DataFrame({
        "Metric": [
            "Overall Risk Status",
            "Annualised Volatility",
            "95% VaR",
        ],
        "Value": [
            "PASS",
            0.05,
            -0.01,
        ],
    })

    limit_results = pd.DataFrame({
        "Metric": [
            "Volatility",
            "VaR 95%",
        ],
        "Actual": [
            0.05,
            -0.01,
        ],
        "Warning": [
            0.07,
            -0.015,
        ],
        "Limit": [
            0.08,
            -0.02,
        ],
        "Status": [
            "PASS",
            "PASS",
        ],
    })

    factor_dashboard = pd.DataFrame({
        "Factor": [
            "Equity",
            "Rates",
            "Credit",
        ],
        "Factor Exposure": [
            0.30,
            0.10,
            0.05,
        ],
    })

    stress_dashboard = pd.DataFrame({
        "Scenario": [
            "Equity Shock",
            "Credit Shock",
        ],
        "Portfolio Return": [
            -0.06,
            -0.02,
        ],
        "Portfolio P&L": [
            -6000,
            -2000,
        ],
    })

    summary = {
        "Overall Status": "PASS",
        "Number of Breaches": 0,
        "Number of Warnings": 0,
    }

    output_path = tmp_path / "risk_report.xlsx"

    result = create_risk_monitoring_excel(
        dashboard,
        limit_results,
        factor_dashboard,
        stress_dashboard,
        summary,
        output_path,
    )

    assert result.exists()

    workbook = load_workbook(result)

    assert "Risk Monitor" in workbook.sheetnames

    worksheet = workbook["Risk Monitor"]

    assert worksheet["A1"].value == "PORTFOLIO RISK MONITOR"
    assert worksheet["A3"].value == "Overall Risk Status"
    assert worksheet["B3"].value == "PASS"