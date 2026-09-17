from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


def write_executive_summary(ws, dashboard, summary):
    ws["A1"] = "PORTFOLIO RISK MONITOR"
    ws["A1"].font = Font(bold=True, size=16)

    ws["A3"] = "Overall Risk Status"
    ws["B3"] = summary["Overall Status"]

    ws["A4"] = "Risk Limit Breaches"
    ws["B4"] = summary["Number of Breaches"]

    ws["A5"] = "Risk Limit Warnings"
    ws["B5"] = summary["Number of Warnings"]

    ws["A7"] = "Risk Metric"
    ws["B7"] = "Value"

    for cell in ws[7]:
        cell.font = Font(bold=True)

    row = 8

    for _, record in dashboard.iterrows():
        ws.cell(row=row, column=1, value=record["Metric"])
        ws.cell(row=row, column=2, value=record["Value"])
        row += 1

    return row


def write_limit_section(ws, limit_results, start_row):
    ws.cell(
        row=start_row,
        column=1,
        value="RISK LIMIT MONITORING",
    ).font = Font(bold=True, size=14)

    header_row = start_row + 2

    headers = [
        "Metric",
        "Actual",
        "Warning",
        "Limit",
        "Status",
        "Limit Utilisation",
    ]

    for column, header in enumerate(headers, start=1):
        cell = ws.cell(
            row=header_row,
            column=column,
            value=header,
        )
        cell.font = Font(bold=True)

    row = header_row + 1

    for _, record in limit_results.iterrows():
        ws.cell(row=row, column=1, value=record["Metric"])
        ws.cell(row=row, column=2, value=record["Actual"])
        ws.cell(row=row, column=3, value=record["Warning"])
        ws.cell(row=row, column=4, value=record["Limit"])
        ws.cell(row=row, column=5, value=record["Status"])
        ws.cell(
            row=row,
            column=6,
            value=record["Limit Utilisation"],
        )
        row += 1

    return row


def write_factor_section(ws, factor_dashboard, start_row):
    ws.cell(
        row=start_row,
        column=1,
        value="FACTOR RISK",
    ).font = Font(bold=True, size=14)

    header_row = start_row + 2

    headers = list(factor_dashboard.columns)

    for column, header in enumerate(headers, start=1):
        cell = ws.cell(
            row=header_row,
            column=column,
            value=header,
        )
        cell.font = Font(bold=True)

    row = header_row + 1

    for _, record in factor_dashboard.iterrows():
        for column, value in enumerate(record, start=1):
            ws.cell(
                row=row,
                column=column,
                value=value,
            )
        row += 1

    return row


def write_stress_section(ws, stress_dashboard, start_row):
    ws.cell(
        row=start_row,
        column=1,
        value="STRESS TESTING",
    ).font = Font(bold=True, size=14)

    header_row = start_row + 2

    headers = list(stress_dashboard.columns)

    for column, header in enumerate(headers, start=1):
        cell = ws.cell(
            row=header_row,
            column=column,
            value=header,
        )
        cell.font = Font(bold=True)

    row = header_row + 1

    for _, record in stress_dashboard.iterrows():
        for column, value in enumerate(record, start=1):
            ws.cell(
                row=row,
                column=column,
                value=value,
            )
        row += 1

    return row


def format_risk_monitoring_sheet(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="center"
            )

    widths = {
        1: 32,
        2: 18,
        3: 18,
        4: 18,
        5: 18,
        6: 20,
    }

    for column, width in widths.items():
        ws.column_dimensions[
            get_column_letter(column)
        ].width = width

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00%"

    ws.freeze_panes = "A8"


def add_limit_conditional_formatting(ws):
    for row in range(10, 30):
        ws.conditional_formatting.add(
            f"F{row}",
            ColorScaleRule(
                start_type="min",
                start_color="FFFFFF",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFFF00",
                end_type="max",
                end_color="FF0000",
            ),
        )


def create_risk_monitoring_excel(
    dashboard,
    limit_results,
    factor_dashboard,
    stress_dashboard,
    summary,
    output_path="reports/risk_monitoring_report.xlsx",
):
    output_path = Path(output_path)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    ws = workbook.active
    ws.title = "Risk Monitor"

    write_executive_summary(
        ws,
        dashboard,
        summary,
    )

    limit_results = limit_results.copy()

    if "Limit Utilisation" not in limit_results.columns:
        limit_results["Limit Utilisation"] = (
            limit_results["Actual"].abs()
            / limit_results["Limit"].abs()
        )

    next_row = write_limit_section(
        ws,
        limit_results,
        20,
    )

    next_row = write_factor_section(
        ws,
        factor_dashboard,
        next_row + 2,
    )

    write_stress_section(
        ws,
        stress_dashboard,
        next_row + 2,
    )

    format_risk_monitoring_sheet(ws)
    add_limit_conditional_formatting(ws)

    workbook.save(output_path)

    return output_path